"""Synthetic catalog procedures. Inputs and snapshots are read-only; no mutations.

This deliberately conservative text extractor handles the documented fixture
formats. Image-only, unknown and contradictory documents require clarification.
It never reads evaluator answer keys or treats filenames as business evidence.
"""
from pathlib import Path
from decimal import Decimal, InvalidOperation
import csv
import json
import os
import re
import sqlite3
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader

INPUTS = Path(os.environ.get("WORKFLOW_INPUTS", "/workspace/inputs"))
REFERENCES = Path(os.environ.get("WORKFLOW_REFERENCES", "/workspace/references"))
OUTPUTS = Path(os.environ.get("WORKFLOW_OUTPUTS", "/workspace/output"))
PACKAGE = Path(os.environ.get("WORKFLOW_PACKAGE", "/workspace/package"))
PREVIOUS = Path(os.environ.get("WORKFLOW_PRECEDING", "/workspace/preceding"))

ALIASES = {
    "invoice_number": ["invoice number", "invoice no", "hesab faktura"],
    "supplier_id": ["supplier id", "təchizatçı", "techizatci"],
    "po_number": ["purchase order", "po number", "sifariş", "sifaris"],
    "currency": ["currency", "valyuta"], "bank": ["bank", "iban"],
    "sku": ["sku", "məhsul", "mehsul"], "quantity": ["quantity", "miqdar"],
    "unit_price": ["unit price", "vahid qiymət", "vahid qiymet"],
    "total": ["total", "cəmi", "cemi"],
    "application_id": ["application id", "müraciət", "muraciet"],
    "company_id": ["company id", "voen", "vöen"],
    "company_name": ["company name", "şirkət", "sirket"],
    "representative": ["representative", "nümayəndə", "numayende"],
    "representative_id": ["representative id", "id number", "şəxsiyyət", "sexsiyyet"],
    "registration": ["registration", "qeydiyyat"], "stamp": ["stamp", "möhür", "mohur"],
    "authority": ["authority", "səlahiyyət", "selahiyyet"],
    "signer_role": ["signer role", "imzalayan"], "valid_until": ["valid until", "etibarlıdır"],
    "complaint_id": ["complaint id", "şikayət", "sikayet"],
    "account_id": ["account id", "abunəçi", "abuneci"],
    "billing_period": ["billing period", "dövr", "dovr"],
}


class MissingEvidence(ValueError):
    def __init__(self, source):
        self.source = source
        super().__init__(f"Required reference snapshot {source} is missing")


def text_of(file):
    if file.suffix.lower() == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(file).pages)
    if file.suffix.lower() == ".docx":
        document = Document(file)
        return "\n".join([p.text for p in document.paragraphs] +
                         ["\n".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows])
    if file.suffix.lower() in [".txt", ".md"]:
        return file.read_text(encoding="utf-8")
    return ""


def extract(file):
    text = text_of(file)
    fields = {}
    conflicts = []
    for key, labels in ALIASES.items():
        values = []
        for label in labels:
            values.extend(re.findall(r"(?im)^\s*" + re.escape(label) + r"\s*[:=]\s*([^\n\r]+)", text))
        values = list(dict.fromkeys(value.strip() for value in values if value.strip()))
        if len(values) == 1:
            fields[key] = values[0]
        elif len(values) > 1:
            conflicts.append(key)
    return fields, conflicts, text


def documents(prefix=None):
    files = sorted(p for p in INPUTS.rglob("*") if p.is_file() and
                   p.suffix.lower() in [".pdf", ".docx", ".txt", ".md"])
    if prefix:
        named = [p for p in files if p.name.lower().startswith(prefix)]
        if named:
            return named
        discriminators = {"invoice": "total", "application": "application_id", "extract": "registration", "authority": "authority", "identification": "valid_until", "complaint": "complaint_id"}
        key = discriminators.get(prefix)
        files = [p for p in files if key in extract(p)[0]] if key else []
    return files


def snapshot(name):
    candidates = list(REFERENCES.rglob(name)) + list(INPUTS.rglob(name))
    if not candidates:
        raise MissingEvidence(name)
    return candidates[0]


def db(name):
    conn = sqlite3.connect("file:" + str(snapshot(name)) + "?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def money(value):
    # Preserve input precision for exact price and tolerance comparisons. Only
    # final displayed/queued amounts are rounded to currency precision.
    number = Decimal(str(value).replace(",", "."))
    if not number.is_finite():
        raise InvalidOperation("Financial values must be finite")
    return number


def issue(code, message, evidence, severity="block"):
    return {"code": code, "message": message, "severity": severity,
            "evidence": [{"source": evidence, "location": code, "quote": message}]}


def missing_reference(source):
    return issue("missing_reference", f"Provide the required {source} reference snapshot.", source, "clarification")


def finish(key, extracted, findings, queue=None, **extra):
    disposition = "clarification" if any(f["severity"] == "clarification" for f in findings) else \
        "escalated" if any(f["severity"] == "escalate" for f in findings) else \
        "blocked" if findings else "ready"
    return {"businessKey": key, "disposition": disposition, "ready": disposition == "ready",
            "findings": findings, "extracted": extracted, "queueRecord": queue or {},
            "synthetic": True, **extra}


def invoice():
    files = documents("invoice") or [p for p in documents() if p.suffix.lower() == ".pdf"]
    if not files:
        return finish("unidentified-invoice", {}, [issue("missing_invoice", "Upload the invoice PDF.", "inputs", "clarification")])
    fields, conflicts, text = extract(files[0])
    findings = [issue("conflicting_field", f"Conflicting {key} values require clarification.", files[0].name, "clarification") for key in conflicts]
    required = ["invoice_number", "supplier_id", "po_number", "currency", "bank", "sku", "quantity", "unit_price", "total"]
    for key in required:
        if key not in fields:
            findings.append(issue("missing_" + key, "Provide a readable " + key.replace("_", " ") + ".", files[0].name, "clarification"))
    business_key = fields.get("supplier_id", "unknown") + ":" + fields.get("invoice_number", "unknown")
    if any(key not in fields for key in required) or conflicts:
        return finish(business_key, fields, findings)
    try:
        quantity, price, declared_total = money(fields["quantity"]), money(fields["unit_price"]), money(fields["total"])
        if quantity <= 0:
            raise InvalidOperation()
    except (InvalidOperation, ValueError):
        return finish(business_key, fields, [issue("invalid_numeric", "Quantity must be positive; quantity, unit price and total must be readable finite numbers.", files[0].name, "clarification")])
    for source in ["purchase-orders.xlsx", "supplier-records.sqlite"]:
        try:
            snapshot(source)
        except MissingEvidence as error:
            findings.append(missing_reference(error.source))
    if findings:
        return finish(business_key, fields, findings)
    purchase_orders = load_workbook(snapshot("purchase-orders.xlsx"), read_only=True, data_only=True)
    sheet = purchase_orders["Orders"]
    rows = list(sheet.values)
    orders = [dict(zip(rows[0], row)) for row in rows[1:] if row[0]]
    conn = db("supplier-records.sqlite")
    try:
        supplier = conn.execute("SELECT * FROM suppliers WHERE supplier_id=?", (fields["supplier_id"],)).fetchone()
        if not supplier:
            findings.append(issue("unknown_supplier", "Supplier is absent from the approved supplier snapshot.", "supplier-records.sqlite:suppliers"))
        elif fields["bank"] != supplier["bank"]:
            findings.append(issue("bank_mismatch", "Invoice bank details do not match the approved supplier.", "supplier-records.sqlite:suppliers"))
        if fields["currency"] != "AZN":
            findings.append(issue("currency_mismatch", "The fictional policy approves AZN only.", "policy.docx:Currency"))
        matched = [row for row in orders if str(row["po_number"]) == fields["po_number"]]
        if not matched:
            findings.append(issue("unknown_po", "Purchase order was not found.", "purchase-orders.xlsx:Orders"))
        else:
            if any(row["supplier_id"] != fields["supplier_id"] for row in matched):
                findings.append(issue("supplier_po_mismatch", "Supplier does not correspond to the purchase order.", "purchase-orders.xlsx:Orders"))
            aliases = [row["sku"] for row in conn.execute("SELECT sku FROM aliases WHERE alias=?", (fields["sku"],))]
            sku = fields["sku"] if any(row["sku"] == fields["sku"] for row in matched) else aliases[0] if len(aliases) == 1 else None
            lines = [row for row in matched if row["sku"] == sku]
            if len(lines) != 1:
                findings.append(issue("ambiguous_sku", "SKU must identify exactly one order line or approved alias.", "purchase-orders.xlsx:Orders; supplier-records.sqlite:aliases", "clarification"))
            else:
                row = lines[0]
                delivered = Decimal(str(conn.execute("SELECT COALESCE(SUM(quantity),0) FROM deliveries WHERE po_number=? AND sku=?", (fields["po_number"], sku)).fetchone()[0]))
                billed = Decimal(str(conn.execute("SELECT COALESCE(SUM(quantity),0) FROM billing WHERE po_number=? AND sku=? AND approved=1", (fields["po_number"], sku)).fetchone()[0]))
                remaining_order = Decimal(str(row["quantity"])) - billed
                remaining_delivery = delivered - billed
                fields.update({"matched_sku": sku, "delivered": str(delivered), "prior_approved_quantity": str(billed)})
                if quantity > remaining_order:
                    findings.append(issue("excess_order_quantity", f"Quantity {quantity} exceeds remaining order {remaining_order}.", "purchase-orders.xlsx:Orders; supplier-records.sqlite:billing"))
                if quantity > remaining_delivery:
                    findings.append(issue("excess_delivery_quantity", f"Quantity {quantity} exceeds remaining delivery {remaining_delivery}.", "supplier-records.sqlite:deliveries,billing"))
                if price != money(row["unit_price"]):
                    findings.append(issue("price_mismatch", f"Unit price {price} differs from agreed {money(row['unit_price'])}.", "purchase-orders.xlsx:Orders"))
        raw_calculated = quantity * price
        calculated = raw_calculated.quantize(Decimal("0.01"))
        tolerance = Decimal("0.01")
        policy_settings = PACKAGE / "settings.json"
        if policy_settings.exists():
            tolerance = Decimal(str(json.loads(policy_settings.read_text()).get("totalTolerance", "0.01")))
        if abs(raw_calculated - declared_total) > tolerance:
            findings.append(issue("total_mismatch", f"Declared total {declared_total} differs from calculated {raw_calculated} beyond tolerance {tolerance} AZN.", "policy.docx:Totals; " + files[0].name))
        if conn.execute("SELECT 1 FROM billing WHERE supplier_id=? AND invoice_number=? AND approved=1", (fields["supplier_id"], fields["invoice_number"])).fetchone():
            findings.append(issue("duplicate_invoice", "This supplier invoice is already in approved billing.", "supplier-records.sqlite:billing"))
        queue = {"supplierId": fields["supplier_id"], "invoiceNumber": fields["invoice_number"], "poNumber": fields["po_number"], "currency": fields["currency"], "amount": float(calculated), "quantity": float(quantity)}
        return finish(business_key, fields, findings, queue, calculatedTotal=float(calculated))
    finally:
        conn.close()
        purchase_orders.close()


def onboarding():
    found = {}
    findings = []
    for name in ["application", "extract", "authority", "identification"]:
        files = documents(name)
        if files:
            fields, conflicts, _ = extract(files[0])
            found[name] = fields
            for key in conflicts:
                findings.append(issue("conflicting_field", f"Conflicting {key} in {name}.", files[0].name, "clarification"))
    app = found.get("application", {})
    company = app.get("company_id")
    key = app.get("application_id", "unidentified-application")
    for name in ["application", "extract", "authority", "identification"]:
        if name not in found:
            findings.append(issue("missing_" + name, f"Provide the representative {'identification document' if name == 'identification' else name + ' document'}.", "inputs", "clarification"))
    required = {"application": ["application_id", "company_id", "company_name", "representative", "signer_role"], "extract": ["company_id", "company_name", "registration"], "authority": ["company_id", "representative", "representative_id", "authority"], "identification": ["representative", "representative_id", "valid_until"]}
    for name, keys in required.items():
        if name in found:
            for field in keys:
                if field not in found[name]:
                    findings.append(issue("missing_" + field, f"Clarify {field} in {name}.", name, "clarification"))
    for name, fields in found.items():
        if fields.get("company_id") and company and fields["company_id"] != company:
            findings.append(issue("company_conflict", f"Company identifier in {name} disagrees with application.", name))
        if fields.get("representative") and app.get("representative") and fields["representative"] != app["representative"]:
            findings.append(issue("representative_conflict", f"Representative in {name} disagrees with application.", name))
    authority, identity, extract_fields = found.get("authority", {}), found.get("identification", {}), found.get("extract", {})
    if authority.get("representative_id") and identity.get("representative_id") and authority["representative_id"] != identity["representative_id"]:
        findings.append(issue("identification_conflict", "Identity document does not match the authorized representative.", "authority; identification"))
    if authority.get("authority") and authority["authority"].lower() != "subscriber contract":
        findings.append(issue("insufficient_authority", "Power of attorney does not authorize a subscriber contract.", "authority"))
    if identity.get("valid_until", "9999") < "2026-10-09":
        findings.append(issue("expired_identification", "Identification has expired as of the synthetic snapshot date.", "identification"))
    if app.get("signer_role") == "head" and app.get("stamp", "no").lower() != "yes":
        findings.append(issue("missing_contract_stamp", "A head-signed contract requires the company stamp.", "procedure.docx:Conditional stamp"))
    # An electronic registration extract has no stamp: this is valid and must not
    # be confused with the separate head-signed contract stamp condition.
    if company:
        try:
            conn = db("registry.sqlite")
        except MissingEvidence as error:
            return finish(key, found, findings + [missing_reference(error.source)], employeeChecks=["Check signature authenticity and original representative authority before approval."], stampRequired=app.get("signer_role") == "head")
        try:
            row = conn.execute("SELECT * FROM registry WHERE company_id=?", (company,)).fetchone()
            if not row:
                findings.append(issue("registry_missing", "Official-register snapshot does not contain the company.", "registry.sqlite:registry", "clarification"))
            elif row["company_name"] != app.get("company_name") or row["company_name"] != extract_fields.get("company_name"):
                findings.append(issue("registry_conflict", "Company name conflicts with the synthetic official-register snapshot.", "registry.sqlite:registry; extract"))
        finally:
            conn.close()
    queue = {"applicationId": key, "companyId": company, "representative": app.get("representative"), "signatureCheck": "employee_required"}
    return finish(key, found, findings, queue, employeeChecks=["Check signature authenticity and original representative authority before approval."], stampRequired=app.get("signer_role") == "head")


def complaints():
    files = documents("complaint")
    if not files:
        return finish("unidentified-complaint", {}, [issue("missing_complaint", "Provide the complaint.", "inputs", "clarification")])
    fields, conflicts, _ = extract(files[0])
    findings = []
    for key in ["complaint_id", "account_id", "billing_period"]:
        if key not in fields:
            findings.append(issue("missing_" + key, f"Provide {key}.", files[0].name, "clarification"))
    for key in conflicts:
        findings.append(issue("conflicting_field", f"Conflicting {key} values.", files[0].name, "clarification"))
    key = fields.get("complaint_id", "unidentified-complaint")
    invoice_files = documents("invoice")
    invoice_fields = extract(invoice_files[0])[0] if invoice_files else {}
    if not invoice_fields.get("total"):
        findings.append(issue("missing_invoice", "Provide a readable bill with its total.", "inputs", "clarification"))
    events_path = next(iter(INPUTS.rglob("events.csv")), None)
    if not events_path:
        return finish(key, fields, findings + [issue("missing_events", "Charge and usage evidence is missing; escalate investigation.", "inputs", "escalate")], creditRecommendation=0)
    with events_path.open(newline="", encoding="utf-8") as file:
        events = list(csv.DictReader(file))
    scoped = [event for event in events if event.get("account_id") == fields.get("account_id") and event.get("billing_period") == fields.get("billing_period")]
    charges = [event for event in scoped if event.get("kind") == "charge"]
    usage = {event["event_id"]: event for event in scoped if event.get("kind") == "usage"}
    seen_transactions = {}
    unique_charges = []
    for event in charges:
        previous = seen_transactions.get(event["transaction_id"])
        if previous and previous != event:
            findings.append(issue("transaction_conflict", "The same transaction ID contains conflicting evidence.", "events.csv", "escalate"))
        elif not previous:
            unique_charges.append(event)
            seen_transactions[event["transaction_id"]] = event
    grouped = {}
    charged_total = Decimal("0")
    for event in unique_charges:
        if event["event_id"] not in usage:
            findings.append(issue("missing_usage_evidence", "Charge has no matching billable usage event: " + event["event_id"], "events.csv", "escalate"))
        try:
            charged_total += money(event["amount"])
            grouped.setdefault(event["event_id"], []).append(event)
        except InvalidOperation:
            findings.append(issue("invalid_charge", "Unreadable charge amount.", "events.csv", "clarification"))
    credit = Decimal("0")
    duplicate_evidence = []
    for event_id, group in grouped.items():
        ordered = sorted(group, key=lambda event: (event["timestamp"], event["transaction_id"]))
        if len(ordered) > 1:
            credit += sum((money(event["amount"]) for event in ordered[1:]), Decimal("0"))
            duplicate_evidence.append({"eventId": event_id, "transactions": [event["transaction_id"] for event in ordered], "credit": float(sum((money(event["amount"]) for event in ordered[1:]), Decimal("0")))})
    if invoice_fields.get("total") and abs(money(invoice_fields["total"]) - charged_total) > Decimal("0.01"):
        findings.append(issue("bill_reconciliation", "Bill total does not reconcile with the charge ledger.", "invoice; events.csv", "escalate"))
    try:
        conn = db("accounts.sqlite")
    except MissingEvidence as error:
        return finish(key, fields, findings + [missing_reference(error.source)], creditRecommendation=float(credit), chargedTotal=float(charged_total), duplicateEvidence=duplicate_evidence, timeline=scoped)
    try:
        account = conn.execute("SELECT * FROM accounts WHERE account_id=?", (fields.get("account_id", ""),)).fetchone()
        if not account:
            findings.append(issue("unknown_account", "Account not found in account snapshot.", "accounts.sqlite:accounts", "clarification"))
        if conn.execute("SELECT 1 FROM resolutions WHERE complaint_id=?", (key,)).fetchone():
            findings.append(issue("duplicate_resolution", "Complaint was already resolved.", "accounts.sqlite:resolutions"))
        payments = [dict(row) for row in conn.execute("SELECT * FROM payments WHERE account_id=?", (fields.get("account_id", ""),))]
    finally:
        conn.close()
    queue = {"complaintId": key, "accountId": fields.get("account_id"), "creditAmount": float(credit), "currency": "AZN", "resolution": "credit_duplicate" if credit else "explain_legitimate"}
    return finish(key, fields, findings, queue, creditRecommendation=float(credit), chargedTotal=float(charged_total), duplicateEvidence=duplicate_evidence, timeline=scoped, payments=payments,
                  recommendation="Credit distinct repeated transactions for the same billable event." if credit else "Explain the supported legitimate charges.")


def workbook(name, title, headers, rows):
    target = OUTPUTS / name
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.append([title])
    sheet.append(["Fictional company data and policy. Employee approval required."])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i, column in enumerate(sheet.columns):
        from openpyxl.utils import get_column_letter
        length = max((len(str(c.value or "")) for c in list(column)[2:]), default=12)
        sheet.column_dimensions[get_column_letter(i + 1)].width = min(max(length + 3, 18), 45)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet["A1"].font = Font(size=16, bold=True)
    sheet["A2"].font = Font(size=10, color="555555")
    for row in sheet.iter_rows(min_row=3):
        required_lines = max((max(1, (len(str(cell.value or "")) + int(sheet.column_dimensions[cell.column_letter].width) - 1) // int(sheet.column_dimensions[cell.column_letter].width)) for cell in row), default=1)
        sheet.row_dimensions[row[0].row].height = 18 * required_lines + 8
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 30
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
    book.save(target)
    return target


def letter(name, title, result, paragraphs):
    target = OUTPUTS / name
    templates = list(REFERENCES.rglob("letter-template.docx")) + list(INPUTS.rglob("letter-template.docx"))
    template = templates[0] if templates else PACKAGE / "templates" / "letter.docx"
    document = Document(template) if template.exists() else Document()
    # Retain template styles and page settings, replacing its instructional body.
    for element in list(document.element.body):
        if not element.tag.endswith("}sectPr"):
            document.element.body.remove(element)
    document.sections[0].top_margin = Inches(.8)
    document.sections[0].bottom_margin = Inches(.8)
    document.add_paragraph(title, "Title")
    document.add_paragraph("Fictional demonstration company. This draft requires employee review and is not sent externally.")
    document.add_paragraph("Business reference " + result["businessKey"])
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        document.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
        for border in document.styles[style_name]._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.save(target)
    return target


def render(kind, result):
    files = []
    if kind == "invoice":
        fields = result["extracted"]
        files.append(workbook("review.xlsx", "Supplier invoice review", ["Business key", "Disposition", "Quantity", "Unit price AZN", "Calculated total AZN", "Findings"], [[result["businessKey"], result["disposition"], fields.get("quantity"), fields.get("unit_price"), result.get("calculatedTotal"), "; ".join(f["code"] for f in result["findings"])]]))
        if result["disposition"] != "ready":
            files.append(letter("correction-letter.docx", "Supplier invoice correction request", result, ["Please correct or clarify the following invoice details before we can submit the accounting record."] + [finding["message"] for finding in result["findings"]]))
    elif kind == "onboarding":
        files.append(workbook("checklist.xlsx", "Corporate subscriber checklist", ["Document", "Provided", "Business key", "Disposition"], [[name, name in result["extracted"], result["businessKey"], result["disposition"]] for name in ["application", "extract", "authority", "identification"]]))
        files.append(letter("application-packet.docx", "Corporate subscriber application packet", result, ["Application disposition " + result["disposition"], "An electronic extract does not require a stamp. Contract stamp requirements depend on the signer."] + result["employeeChecks"]))
        if result["findings"]:
            files.append(letter("missing-information.docx", "Corporate subscriber information request", result, [finding["message"] for finding in result["findings"]]))
    else:
        files.append(workbook("reconciliation.xlsx", "Billing complaint reconciliation", ["Business key", "Disposition", "Charged AZN", "Credit AZN", "Recommendation"], [[result["businessKey"], result["disposition"], result.get("chargedTotal"), result.get("creditRecommendation", 0), result.get("recommendation", "Clarify or escalate the missing evidence.")]]))
        paragraphs = ["Investigation disposition " + result["disposition"], result.get("recommendation", "Further evidence is required."), "Recommended credit AZN " + str(result.get("creditRecommendation", 0))] + [finding["message"] for finding in result["findings"]]
        for duplicate in result.get("duplicateEvidence", []):
            paragraphs.append("Event " + duplicate["eventId"] + " was billed by distinct transactions " + ", ".join(duplicate["transactions"]) + ".")
        files.append(letter("response.docx", "Billing complaint investigation response", result, paragraphs))
        (OUTPUTS / "timeline.json").write_text(json.dumps(result.get("timeline", []), ensure_ascii=False, indent=2))
        files.append(OUTPUTS / "timeline.json")
    return [{"path": p.name, "name": p.name, "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if p.suffix == ".xlsx" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document" if p.suffix == ".docx" else "application/json"} for p in files]


def run(kind):
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    try:
        result = {"invoice": invoice, "onboarding": onboarding, "complaints": complaints}[kind]()
    except MissingEvidence as error:
        result = finish("unidentified-" + kind, {}, [missing_reference(error.source)])
    artifacts = render(kind, result)
    (OUTPUTS / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
    (OUTPUTS / "manifest.json").write_text(json.dumps({"stepId": os.environ.get("WORKFLOW_STEP_ID", "validate"), "result": result, "artifacts": artifacts}, ensure_ascii=False, indent=2))
    print(json.dumps({"disposition": result["disposition"], "findings": len(result["findings"]), "manifest": "manifest.json"}))


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        run(sys.argv[1])
